#!/usr/bin/env python3
"""Gera data/data.json a partir dos dois Excel do SAP.

Uso:
  python3 scripts/preparar_dados.py Infos_modelo.xlsx base.xlsx [--corte "25/10/2026 10:29"] [--corte-seg "25/10/2026 13:11"]

Requer: pip install openpyxl pandas
  Infos_modelo.xlsx  aba "Infos modelo": metas e volumes por Gerência Agrupada x Rota x pilar
  base.xlsx          aba "base": Gerência, Supervisão, Rota, Segmentação e volume
Regras de montagem:
  - canal: Segmentada Mini/Mod -> Segmentada; rota em dois segmentos fica com o de maior volume;
  - rota do modelo sem segmentação na base vai para "semCanal" (fora do painel);
  - gerência e supervisão vêm da base (Gerência Comercial); nomes de SV normalizados ("Sv. G3. 2" -> "Sv. G3.2").
Quando a API estiver pronta, este script deixa de ser necessário: troque só o carregamento de dados.
"""
import argparse, json, math, re, pathlib
import pandas as pd
from openpyxl import load_workbook

ap = argparse.ArgumentParser()
ap.add_argument('infos'); ap.add_argument('base')
ap.add_argument('--corte', default='25/10/2026 10:29')
ap.add_argument('--corte-seg', default='25/10/2026 13:11')
ap.add_argument('--saida', default=str(pathlib.Path(__file__).resolve().parent.parent / 'data' / 'data.json'))
a = ap.parse_args()

PRODS = ['Powerade', 'Refrigerantes Individual', 'Refrigerantes Retornavel', 'Sabores Tt', 'Sucos']
MEDS = ['Meta', 'Real', 'WF Online', 'Hoje', 'MetaDia']

ws = load_workbook(a.infos, data_only=True)['Infos modelo']
recs = []
for r in range(4, ws.max_row + 1):
    d = {'gc': ws.cell(r, 1).value, 'rota': str(ws.cell(r, 2).value)}
    for i, p in enumerate(PRODS):
        for j, m in enumerate(MEDS):
            d[f'{p}|{m}'] = ws.cell(r, 3 + i * 5 + j).value
    recs.append(d)
df = pd.DataFrame(recs)

wb = load_workbook(a.base, data_only=True)['base']
b = pd.DataFrame([(wb.cell(r, 1).value, wb.cell(r, 2).value, str(wb.cell(r, 3).value), wb.cell(r, 4).value, wb.cell(r, 5).value)
                  for r in range(4, wb.max_row + 1)], columns=['gc', 'sv', 'rota', 'seg', 'vol'])
b = b[b.rota != '--']  # linhas de total

nn = lambda x: None if x is None or (isinstance(x, float) and math.isnan(x)) else round(float(x), 4)
svn = lambda s: re.sub(r'\.\s+(\d)', r'.\1', str(s).strip())
gcn = lambda g: g.replace('Gc. ', '').replace('_', '/')

dup, esc = [], {}
for rota, g in b.groupby('rota'):
    g = g.sort_values('vol', ascending=False); top = g.iloc[0]
    esc[rota] = (top.seg, svn(top.sv), gcn(top.gc))
    if len(g) > 1:
        dup.append(dict(rota=rota, escolhida=top.seg, opcoes=[dict(seg=r.seg, vol=round(float(r.vol), 1)) for r in g.itertuples()]))

rotas, semcanal = [], []
for _, r in df.iterrows():
    if r.rota not in esc:
        semcanal.append(dict(gc=gcn(r.gc), rota=r.rota)); continue
    s, sv, gc = esc[r.rota]
    rotas.append(dict(gc=gc, sv=sv, rota=r.rota, canal='Mista' if s == 'Mista' else 'Segmentada',
                      sub=None if s == 'Mista' else s.replace('Segmentada ', ''),
                      v={p: [nn(r[f'{p}|{m}']) for m in MEDS] for p in PRODS}))

out = dict(meta=dict(mes=10, ano=2026, corte=a.corte, corteSeg=a.corte_seg,
                     fonte=f'SAP Analytics Cloud · metas e volumes {a.corte.split()[-1]} · segmentação {a.corte_seg.split()[-1]}',
                     fonteLonga=f'SAP Analytics Cloud · modelo "Infos modelo" (metas e volumes, Data 202610, {a.corte.split()[-1]}) e modelo "base" (gerência, supervisão e segmentação, {a.corte_seg.split()[-1]})',
                     medidas=MEDS, produtos=PRODS, totalRotasModelo=len(df)),
           rotas=rotas, semCanal=semcanal, duplicadas=dup, segSemInfos=sorted(set(b.rota) - set(df.rota)))
pathlib.Path(a.saida).write_text(json.dumps(out, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
print(f'{len(rotas)} rotas com canal, {len(semcanal)} sem canal, {len(dup)} em dois segmentos -> {a.saida}')
